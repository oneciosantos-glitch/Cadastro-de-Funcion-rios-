"""Exportacao dos dados para Excel (.xlsx) ou CSV."""

import csv
import io
from datetime import date

import calculos as cal

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    TEM_OPENPYXL = True
except ImportError:
    TEM_OPENPYXL = False


def _linhas_exportacao(registros, hoje):
    """Gera lista de dicts prontos para exportacao."""
    linhas = []
    for r in registros:
        adm = cal.parse_data(r["admissao"])
        f = cal.calcular_ferias(adm, hoje)
        exp_sit = "-"
        exp_prazo = "-"
        exp_fim = "-"
        if r.get("experiencia_dias"):
            e = cal.contrato_experiencia(adm, r["experiencia_dias"], hoje)
            exp_prazo = f"{e['prazo_dias']}d ({e['etapas']})"
            exp_fim = cal.fmt(e["fim"])
            exp_sit = e["situacao"]

        linhas.append({
            "Matricula": r["matricula"],
            "Funcionario": r["nome"],
            "RG": r.get("rg", ""),
            "CPF": r.get("cpf", ""),
            "Telefone": r.get("telefone", ""),
            "Admissao": cal.fmt(adm),
            "Demissao": cal.fmt(cal.parse_data(r["demissao"])) if r.get("demissao") else "",
            "Loja": r.get("loja", ""),
            "Situacao": r.get("situacao", ""),
            "Cargo": r.get("cargo", ""),
            "Experiencia": exp_prazo,
            "Fim experiencia": exp_fim,
            "Status experiencia": exp_sit,
            "Regra ferias": f["regra"],
            "Periodo ferias": f"{cal.fmt(f['inicio_periodo'])} a {cal.fmt(f['fim_periodo'])}",
            "Liberacao ferias": cal.fmt(f["data_liberacao"]),
            "Meses cumpridos": f["progresso"],
            "Dias proporcionais": f["dias_proporcionais"],
            "Limite gozo": cal.fmt(f["limite_gozo"]),
            "Ferias": f["situacao"],
            "Foto": "Sim" if r.get("foto") else "Nao",
            "Observacao": r.get("observacao", ""),
        })
    return linhas


def exportar_bytes(registros, hoje=None):
    """Retorna (conteudo_bytes, nome_arquivo)."""
    hoje = hoje or date.today()
    linhas = _linhas_exportacao(registros, hoje)
    if TEM_OPENPYXL:
        return _exportar_xlsx(linhas, hoje)
    return _exportar_csv(linhas, hoje)


def _exportar_xlsx(linhas, hoje):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Funcionarios"

    cabecalho = list(linhas[0].keys()) if linhas else []
    ws.append(cabecalho)

    amarelo = PatternFill(start_color="FDF1DC", end_color="FDF1DC",
                          fill_type="solid")
    verde = PatternFill(start_color="E9F7EE", end_color="E9F7EE",
                        fill_type="solid")
    negrito = Font(bold=True)

    for cell in ws[1]:
        cell.font = negrito
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for i, linha in enumerate(linhas, start=2):
        ws.append(list(linha.values()))
        exp_sit = str(linha.get("Status experiencia", ""))
        fer = str(linha.get("Ferias", ""))
        if "atencao" in exp_sit.lower() or "vence" in exp_sit.lower():
            for cell in ws[i]:
                cell.fill = amarelo
        elif "LIBERADA" in fer:
            for cell in ws[i]:
                cell.fill = verde

    for col_idx, _ in enumerate(cabecalho, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    nome = f"funcionarios_{hoje:%Y-%m-%d}.xlsx"
    return buf.getvalue(), nome


def _exportar_csv(linhas, hoje):
    buf = io.StringIO()
    if linhas:
        writer = csv.DictWriter(buf, fieldnames=list(linhas[0].keys()),
                                delimiter=";", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(linhas)
    nome = f"funcionarios_{hoje:%Y-%m-%d}.csv"
    return buf.getvalue().encode("utf-8-sig"), nome
